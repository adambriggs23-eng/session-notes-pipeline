#!/usr/bin/env python3
"""
05_draft_note.py  (Windows version)

Uses a LOCAL LLM (via Ollama, fully offline) to draft a DAP progress note
from a speaker-labelled transcript, and writes it into a Word document
using the same structure as DAP_Note_Template.docx.

================================================================================
IMPORTANT - READ BEFORE USING
================================================================================
This produces a DRAFT ONLY. The output file is named with "_DRAFT" and a
banner is added at the top of the document as a reminder.

A local LLM is much smaller and less capable than a cloud model. It WILL:
  - miss clinical nuance
  - sometimes phrase things generically or awkwardly
  - occasionally misread the transcript

You must review and rewrite every section before this becomes a real
clinical note, especially:
  - Safety / Risk Assessment section - verify against the transcript yourself,
    do not rely on the model to catch risk indicators
  - Assessment section - clinical interpretation, prognosis, and diagnostic
    considerations require your professional judgment, not the model's
  - Any client quotes - verify these against the transcript for accuracy
================================================================================

Usage:
    python 05_draft_note.py <name>

    Where <name> matches the basename from earlier steps, e.g. for
    "session1.m4a" run:
        python 05_draft_note.py session1

Expects:
    transcripts\\<name>_speaker_transcript.txt
        (after you've done the Clinician/Client find-and-replace)

Output:
    notes\\<name>_DRAFT.docx

Environment variables:
    OLLAMA_MODEL        - Model to use (default: llama3.1:8b)
    OLLAMA_TIMEOUT      - Timeout in seconds (default: 1200 = 20 min)
"""

import sys
import os
import json
import re
import urllib.request
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PIPELINE_DIR = Path(os.environ.get("PIPELINE_DIR", Path.home() / "session-pipeline"))
TRANSCRIPTS_DIR = PIPELINE_DIR / "transcripts"
NOTES_DIR = PIPELINE_DIR / "notes"

OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.1:8b")
OLLAMA_TIMEOUT = int(os.environ.get("OLLAMA_TIMEOUT", "1200"))  # 20 minutes default

# ------------------------------------------------------------------------------
# Prompts - one per section. Smaller models do better with focused,
# single-purpose prompts than one giant "fill the whole template" prompt.
# ------------------------------------------------------------------------------

DATA_PROMPT = """You are assisting a mental health clinician by drafting the
DATA section of a DAP (Data-Assessment-Plan) clinical progress note, based on
the session transcript below.

Write in objective, professional clinical language. Use the client's own
words in quotation marks where they describe their experience. Do NOT invent
information that is not in the transcript - if something wasn't discussed,
write "Not discussed in session" for that field.

Respond ONLY with valid JSON, no other text, in exactly this structure:

{{
  "presenting_concerns": "string - what the client identified as their focus for this session, in their words and clinically",
  "client_report": ["bullet point", "bullet point", "..."],
  "clinical_observations": ["bullet point about appearance/behaviour/mood/affect/speech/thought process etc, ONLY based on what is evident from the transcript", "..."],
  "safety_risk_status": "string - any mention of suicidal ideation, homicidal ideation, self-harm, or safety concerns. If none were mentioned, write: 'No suicidal or homicidal ideation reported or observable. No self-harm mentioned.'"
}}

TRANSCRIPT:
{transcript}
"""

ASSESSMENT_PROMPT = """You are assisting a mental health clinician by drafting
the ASSESSMENT section of a DAP clinical progress note, based on the session
transcript below.

This is your DRAFT attempt at clinical interpretation - the clinician will
review and revise this. Be conservative: where the transcript doesn't give
enough information to assess something, say so explicitly rather than
guessing.

Respond ONLY with valid JSON, no other text, in exactly this structure:

{{
  "treatment_goals_addressed": "string - any treatment goals explicitly mentioned in the transcript. If none mentioned, write 'Not explicitly referenced in this session - link to treatment plan goals for context.'",
  "progress_toward_goals": ["bullet point describing any evidence of progress, setbacks, or change mentioned by the client, with supporting detail from the transcript", "..."],
  "client_response_to_interventions": ["bullet point describing how the client engaged with any techniques, suggestions, or exercises discussed", "..."],
  "clinical_interpretation": "string - a draft interpretation based ONLY on what's observable in the transcript (e.g. engagement level, stated progress/setbacks, compliance with prior homework). Do NOT diagnose or predict outcomes. Clinician will add diagnostic/prognostic judgment."
}}

TRANSCRIPT:
{transcript}
"""

PLAN_PROMPT = """You are assisting a mental health clinician by drafting the
PLAN section of a DAP clinical progress note, based on the session transcript
below.

Respond ONLY with valid JSON, no other text, in exactly this structure:

{{
  "interventions_used": ["bullet point naming a specific technique/intervention discussed or used in session, if identifiable from the transcript", "..."],
  "homework_assigned": ["bullet point describing any task, exercise, or between-session activity the client agreed to try", "..."],
  "referrals_coordination": "string - any mention of referrals, other providers, or coordination of care. If none, write 'None discussed.'",
  "safety_plan_updates": "string - any safety planning discussed. If none, write 'Not applicable - no safety concerns raised.'",
  "next_session": "string - any mention of the next appointment date/time or focus. If none, write 'To be scheduled.'"
}}

TRANSCRIPT:
{transcript}
"""


def call_ollama(prompt: str) -> dict:
    """Call local Ollama server and parse JSON response with robust error handling."""
    payload = json.dumps({
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {"temperature": 0.2},
    }).encode("utf-8")

    req = urllib.request.Request(
        OLLAMA_URL, data=payload, headers={"Content-Type": "application/json"}
    )

    try:
        with urllib.request.urlopen(req, timeout=OLLAMA_TIMEOUT) as resp:
            result = json.loads(resp.read().decode("utf-8"))
    except urllib.error.URLError as e:
        print(f"\nERROR: could not reach Ollama at {OLLAMA_URL}")
        print("Is Ollama installed and running? Try:")
        print(f"  ollama run {OLLAMA_MODEL}")
        print(f"\nDetails: {e}")
        sys.exit(1)
    except urllib.error.HTTPError as e:
        print(f"\nERROR: Ollama HTTP error: {e.code} {e.reason}")
        sys.exit(1)
    except Exception as e:
        print(f"\nERROR: Unexpected error connecting to Ollama: {e}")
        sys.exit(1)

    raw_text = result.get("response", "")
    
    # Try parsing as-is first
    try:
        return json.loads(raw_text)
    except json.JSONDecodeError:
        pass
    
    # Try using JSONDecoder.raw_decode() to extract valid JSON object
    try:
        decoder = json.JSONDecoder()
        obj, idx = decoder.raw_decode(raw_text)
        return obj
    except json.JSONDecodeError:
        pass
    
    # Last resort: regex extraction
    match = re.search(r"\{.*\}", raw_text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except json.JSONDecodeError:
            pass
    
    print("\nWARNING: model output was not valid JSON. Raw output was:")
    print(raw_text[:1000])
    return {}


def add_heading(doc, text):
    """Add a formatted heading to the document."""
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)


def add_field(doc, label, value, hint=None):
    """Add a labeled field to the document with optional hint text."""
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True
    if hint:
        hint_run = p.add_run("  " + hint)
        hint_run.italic = True
        hint_run.font.size = Pt(9)
        hint_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    if isinstance(value, list):
        if not value:
            doc.add_paragraph("Not discussed in session.")
        for item in value:
            doc.add_paragraph(str(item), style="List Bullet")
    else:
        doc.add_paragraph(str(value) if value else "Not discussed in session.")


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <name>")
        print(f"  Example: {sys.argv[0]} session1")
        sys.exit(1)

    name = Path(sys.argv[1]).stem  # Strip extension if provided
    transcript_path = TRANSCRIPTS_DIR / f"{name}_speaker_transcript.txt"

    if not transcript_path.exists():
        print(f"Missing: {transcript_path}")
        print("Run the transcription/diarization/merge steps first, and make")
        print("sure you've done the SPEAKER_00/01 -> Clinician/Client rename.")
        sys.exit(1)

    # Create notes directory at start
    NOTES_DIR.mkdir(parents=True, exist_ok=True)

    transcript = transcript_path.read_text(encoding="utf-8")

    if "Clinician" not in transcript and "Client" not in transcript:
        print("NOTE: transcript still uses SPEAKER_00/SPEAKER_01 labels.")
        print("Consider renaming these to Clinician/Client first for a better draft.")
        print("Continuing anyway...\n")

    print(f"Drafting note for: {name}")
    print(f"Using local model: {OLLAMA_MODEL}")
    print(f"Timeout: {OLLAMA_TIMEOUT} seconds\n")

    print("[1/3] Drafting Data section...")
    data = call_ollama(DATA_PROMPT.format(transcript=transcript))

    print("[2/3] Drafting Assessment section...")
    assessment = call_ollama(ASSESSMENT_PROMPT.format(transcript=transcript))

    print("[3/3] Drafting Plan section...")
    plan = call_ollama(PLAN_PROMPT.format(transcript=transcript))

    # ---- Build the document --------------------------------------------------
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    # Fix a minor schema quirk in python-docx's default template (missing zoom percent)
    settings = doc.settings.element
    for zoom in settings.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom"
    ):
        zoom.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent",
            "100",
        )

    title = doc.add_heading("DAP Progress Note \u2013 DRAFT", level=1)

    banner = doc.add_paragraph()
    banner_run = banner.add_run(
        "AI-GENERATED DRAFT \u2014 NOT A FINAL CLINICAL NOTE. "
        "Generated offline from session transcript. Review, verify, and "
        "rewrite all sections before filing, especially Safety/Risk and "
        "Assessment. Verify all quotes against the transcript. "
        "This is a starting point, not a substitute for your clinical judgment."
    )
    banner_run.bold = True
    banner_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    banner_run.font.size = Pt(10)

    # Session info placeholders
    doc.add_paragraph()
    info = doc.add_paragraph()
    info_run = info.add_run(
        "Client Name/ID: ___________    Date of Session: ___________    "
        "Session #: ___________"
    )
    info_run.font.size = Pt(10)

    # ---- DATA ----
    add_heading(doc, "D \u2013 Data")
    add_field(doc, "Presenting concerns / session focus:",
              data.get("presenting_concerns", ""))
    add_field(doc, "Client report (subjective):",
              data.get("client_report", []))
    add_field(doc, "Clinical observations (objective):",
              data.get("clinical_observations", []))
    add_field(doc, "Safety / risk status:",
              data.get("safety_risk_status", ""),
              hint="(CLINICIAN: verify this against transcript independently)")

    # ---- ASSESSMENT ----
    add_heading(doc, "A \u2013 Assessment")
    add_field(doc, "Treatment goal(s) addressed:",
              assessment.get("treatment_goals_addressed", ""))
    add_field(doc, "Progress toward goals:",
              assessment.get("progress_toward_goals", []))
    add_field(doc, "Client response to interventions:",
              assessment.get("client_response_to_interventions", []))
    add_field(doc, "Clinical interpretation / prognosis:",
              assessment.get("clinical_interpretation", ""),
              hint="(CLINICIAN: add diagnostic/prognostic judgment)")

    # ---- PLAN ----
    add_heading(doc, "P \u2013 Plan")
    add_field(doc, "Interventions used this session:",
              plan.get("interventions_used", []))
    add_field(doc, "Homework / between-session tasks assigned:",
              plan.get("homework_assigned", []))
    add_field(doc, "Referrals / coordination of care:",
              plan.get("referrals_coordination", ""))
    add_field(doc, "Safety plan updates / management plan:",
              plan.get("safety_plan_updates", ""))
    add_field(doc, "Plan for next session:",
              plan.get("next_session", ""))

    # Signature placeholder
    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig_run = sig.add_run("Clinician signature: ___________    Date completed: ___________")
    sig_run.bold = True

    out_path = NOTES_DIR / f"{name}_DRAFT.docx"
    doc.save(out_path)

    print(f"\nDraft note written to: {out_path}")
    print("\nREMINDER: Before this becomes a clinical record:")
    print("  1. Read the entire transcript yourself")
    print("  2. Rewrite Safety/Risk section based on YOUR assessment")
    print("  3. Add your clinical judgment to Assessment (diagnosis, prognosis)")
    print("  4. Verify every quoted statement against the transcript")
    print("  5. Rewrite in your own clinical voice")


if __name__ == "__main__":
    main()
